#!/usr/bin/env python3
"""Builds docs/Rasoi.pptx — a walkthrough of the app with the real simulator screenshots.

    python3 scripts/make-deck.py

Screenshots come from docs/screenshots/ (regenerate them with bash scripts/screenshots.sh).
Everything on the slides is a fact from the build: no mockups, no invented numbers.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs/screenshots"
OUT = ROOT / "docs/Rasoi.pptx"

SAFFRON = RGBColor(0xE8, 0xA3, 0x3D)
SAFFRON_DEEP = RGBColor(0xD4, 0x86, 0x2A)
CREAM = RGBColor(0xFB, 0xF6, 0xEC)
INK = RGBColor(0x24, 0x1F, 0x17)
MUTED = RGBColor(0x6B, 0x61, 0x51)
SAGE = RGBColor(0x6F, 0x8F, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)


def blank(prs, background=CREAM):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = background
    return slide


def textbox(slide, left, top, width, height, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.paragraphs[0].alignment = align
    return frame


def write(frame, text, size, color=INK, bold=False, space_after=8, first=False, align=None):
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    paragraph.text = text
    paragraph.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = "Helvetica Neue"
    return paragraph


def accent_bar(slide, color=SAFFRON):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.22), H)  # rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False


def phone(slide, name, left, top=Inches(0.62), height=Inches(6.3)):
    """Places a screenshot, sized by height so the phone aspect ratio is kept."""
    path = SHOTS / f"{name}.png"
    if not path.exists():
        raise SystemExit(f"missing screenshot: {path}")
    return slide.shapes.add_picture(str(path), left, top, height=height)


def title_slide(prs):
    slide = blank(prs, SAFFRON)
    frame = textbox(slide, Inches(0.9), Inches(2.1), Inches(7.4), Inches(3.4))
    write(frame, "Rasoi", 66, WHITE, bold=True, first=True, space_after=4)
    write(frame, "रसोई — “kitchen”", 20, RGBColor(0xFF, 0xF4, 0xDF), space_after=26)
    write(frame, "A private kitchen planner for one family's iPhone.", 26, WHITE, space_after=10)
    write(frame,
          "The pantry drives the plan. The plan drives the list. "
          "The children decide what comes back.",
          18, RGBColor(0xFF, 0xF4, 0xDF), space_after=0)

    phone(slide, "05-today", left=Inches(9.0), top=Inches(0.55), height=Inches(6.4))
    return slide


def statement_slide(prs, kicker, title, lines, bullets=None):
    slide = blank(prs)
    accent_bar(slide)
    frame = textbox(slide, Inches(1.0), Inches(1.15), Inches(11.2), Inches(5.2))
    write(frame, kicker.upper(), 13, SAFFRON_DEEP, bold=True, first=True, space_after=10)
    write(frame, title, 40, INK, bold=True, space_after=20)
    for line in lines:
        write(frame, line, 20, MUTED, space_after=12)
    for bullet in bullets or []:
        write(frame, f"•  {bullet}", 18, INK, space_after=9)
    return slide


def feature_slide(prs, kicker, title, bullets, shots, note=None):
    """Text on the left, one or two screenshots on the right."""
    slide = blank(prs)
    accent_bar(slide)
    frame = textbox(slide, Inches(0.85), Inches(0.95), Inches(6.4), Inches(5.4))
    write(frame, kicker.upper(), 13, SAFFRON_DEEP, bold=True, first=True, space_after=8)
    write(frame, title, 34, INK, bold=True, space_after=18)
    for bullet in bullets:
        write(frame, f"•  {bullet}", 17, INK, space_after=11)
    if note:
        write(frame, note, 14, MUTED, space_after=0)

    if len(shots) == 1:
        phone(slide, shots[0], left=Inches(9.3), height=Inches(6.3))
    else:
        phone(slide, shots[0], left=Inches(7.7), top=Inches(0.95), height=Inches(5.6))
        phone(slide, shots[1], left=Inches(10.5), top=Inches(0.95), height=Inches(5.6))
    return slide


def closing_slide(prs):
    slide = blank(prs, INK)
    frame = textbox(slide, Inches(1.0), Inches(2.3), Inches(11.3), Inches(3.2))
    write(frame, "Built, tested, and quiet about it.", 40, WHITE, bold=True, first=True, space_after=18)
    write(frame,
          "62 automated build tasks · 385 tests green · 22 screens reviewed · "
          "zero third-party code · zero analytics",
          20, RGBColor(0xE8, 0xDE, 0xCB), space_after=14)
    write(frame,
          "Next: install on the iPhone with a free Apple team, then two weeks of real cooking "
          "before Phase 2 (photo scan, AI suggestions, family sync).",
          18, SAFFRON, space_after=0)
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    title_slide(prs)

    statement_slide(
        prs,
        "the problem",
        "Weeknight dinner is a decision nobody wants to make at 6pm",
        ["Four people, two of them children, one vegetarian kitchen, one weekly shop."],
        bullets=[
            "The plan lives in someone's head, so it collapses on a bad Tuesday.",
            "What the children actually eat is remembered by feel, not by record.",
            "Food goes off at the back of the fridge while the list repeats itself.",
            "Every app that would help wants an account, a subscription, and your family's data.",
        ],
    )

    statement_slide(
        prs,
        "how it works",
        "One loop, and it gets better every week",
        ["Pantry  →  Plan  →  List  →  Cook  →  How did it go?  →  a better Plan"],
        bullets=[
            "The planner starts from what is already in the kitchen and what is about to go off.",
            "It ranks meals by what the children actually ate, not by what a magazine says.",
            "Every meal can explain itself: “uses the spinach before it turns”, “25 minutes on a school night”.",
            "Same inputs, same plan — it is deterministic, so a week can be reproduced exactly.",
        ],
    )

    feature_slide(
        prs, "today", "Tonight, in one glance",
        [
            "The dinner, its cuisine, total time and how many it serves.",
            "“Why this” opens the planner's actual reasons.",
            "The rest of the day underneath: breakfast, lunch, snack.",
            "A use-it-up banner when something is about to turn.",
            "Three quick snacks a child can have right now, from the pantry.",
            "Five food-group dots for the day — presence, never portions.",
        ],
        ["05-today"],
    )

    feature_slide(
        prs, "plan", "A whole week in one tap",
        [
            "Seven days × four meals, generated in under a second.",
            "No repeat within three days; no two dinners of the same cuisine running.",
            "Weeknight dinners fit your time cap; the long cook lands at the weekend.",
            "At least two sure things a week, at most two brand-new dishes.",
            "Tap a slot to see why it is there and swap it for a ranked alternative.",
            "Keep, skip or mark eating-out — regenerating never touches those.",
        ],
        ["03-plan", "04-plan-slot"],
    )

    feature_slide(
        prs, "cook", "Cook mode, readable from a metre away",
        [
            "One step per screen, in large type, swipe to move on.",
            "The screen stays awake while you cook.",
            "A servings stepper rescales every quantity — to sensible amounts, not 1.6667 onions.",
        ],
        ["06-cook-ingredients", "07-cook-step"],
    )

    feature_slide(
        prs, "feedback", "Three taps per person, and the app learns",
        [
            "Ate it all · Ate some · Not today — for whoever was at the table.",
            "Skip anyone you didn't get to ask; silence is never counted as a refusal.",
            "Reactions decay: a meal refused four months ago barely counts today.",
            "A dish every child turns down three times running rests for a month, then comes back.",
            "No scores, no streaks, nothing a child could read as a mark against them.",
        ],
        ["08-feedback"],
    )

    feature_slide(
        prs, "pantry", "The kitchen, as it actually is",
        [
            "Fridge, pantry and freezer — the freezer keeps things six times as long.",
            "Type “toor” or “palak”: 185 ingredients, searchable by alias.",
            "Shelf-life dates, soonest first, with plain words: “use today”, “2 days left”.",
            "Mark used or ran out with a swipe; weekly staples stay at zero so the list tops them up.",
        ],
        ["11-pantry", "10-pantry-add"],
    )

    feature_slide(
        prs, "shop", "One list, split the way you shop",
        [
            "Built from the week's plan: quantities summed, units normalised, pantry subtracted.",
            "Grouped by store, then by aisle — dal and spices at the Indian grocery, the rest at the supermarket.",
            "Ticking an item puts it in the pantry with its expiry date already worked out.",
            "“Why?” shows which meals need it. Rebuilding never disturbs what you have already ticked.",
            "Share the whole list as plain text.",
        ],
        ["12-shop", "13-shop-checked"],
    )

    feature_slide(
        prs, "recipes", "51 vegetarian recipes, and room for yours",
        [
            "Indian, Mexican, Italian, Chinese-style, American and Mediterranean.",
            "Filter by meal, cuisine, appliance, quick, favourites or lunchbox.",
            "Every recipe carries a “make it kid-friendly” note.",
            "Ingredients you already have are ticked off against your pantry.",
            "Write your own: allergen flags are worked out from the ingredients, not typed by hand.",
        ],
        ["18-recipes", "19-recipe-detail"],
    )

    feature_slide(
        prs, "insights", "What your own table has taught it",
        [
            "Per child: what goes down well, and what has not lately — counted in words, never scored.",
            "Most-cooked meals and how much variety the last four weeks held.",
            "One gentle weekly hint: “this week is light on fruit”.",
            "Food groups only. No calories, no weight, no diet language, anywhere in the app.",
        ],
        ["20-insights"],
    )

    feature_slide(
        prs, "your household", "Set up once, in under two minutes",
        [
            "Family with dates of birth — ages decide what suits everyone at the table.",
            "Vegetarian is fixed; eggs, dairy, nuts and gluten are yours to set.",
            "Never-suggest list, alias-aware: exclude “palak” and the spinach dishes go too.",
            "Your appliances, your cuisines, your weeknight time cap, your shopping day.",
        ],
        ["16-diet", "15-family"],
    )

    feature_slide(
        prs, "stores & reminders", "Where you shop, and the three reminders",
        [
            "Six stores to start with, reorderable, each claiming the aisles you use it for.",
            "“Find nearby” searches Apple Maps around your zip code — the only network call in the app.",
            "Shopping-day reminder the evening before, cook-tonight, and use-it-up.",
            "One of each a day at most, never between 9pm and 7am.",
        ],
        ["17-stores", "21-notifications"],
    )

    statement_slide(
        prs,
        "privacy",
        "Nothing about your family leaves the phone",
        ["Checked mechanically on every build, not just promised in a settings screen."],
        bullets=[
            "No account, no sync, no analytics, no ads, no third-party code at all.",
            "The one exception: tapping “Find nearby” sends a zip code and five grocery search words to Apple Maps.",
            "scripts/privacy-audit.sh greps the sources, the imports, the dependencies and the built binary — and fails the build on a single hit.",
            "Export gives you a JSON copy of everything; “Delete all data” really deletes it.",
        ],
    )

    statement_slide(
        prs,
        "under the hood",
        "Small, deterministic, and covered by tests",
        ["SwiftUI + SwiftData, iOS 26, one target, zero dependencies."],
        bullets=[
            "Views → view models → pure Swift engines: planner, kid scores, grocery builder, nutrition, units, shelf life.",
            "51 recipes and 185 ingredients ship as JSON data, imported idempotently — data bugs are fixed in data.",
            "385 tests, including one per planner rule, and a plan generated from the real catalog in 0.17s.",
            "22 screens captured and reviewed — five real bugs were found by looking at them, not by the tests.",
        ],
    )

    closing_slide(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides")


if __name__ == "__main__":
    build()
